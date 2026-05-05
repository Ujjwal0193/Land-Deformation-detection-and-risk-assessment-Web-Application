import logging
from pptx import Presentation

def remove_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    if index < len(slides):
        xml_slides.remove(slides[index])

def concenise_text(shape, new_text, font_size=11, bold=False):
    # Overwrite the first frame and remove the rest to "concise" it
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = new_text
    from pptx.util import Pt
    r.font.size = Pt(font_size)
    r.font.bold = bold
    
    # We also want to preserve color if possible, but python-pptx makes this tricky if we clear.
    # To keep it simple, we use a generic white or let the theme handle it.
    from pptx.dml.color import RGBColor
    r.font.color.rgb = RGBColor(255, 255, 255)

def main():
    prs = Presentation('PD_ppt.pptx')
    
    # We will modify text on slides 5, 8, 10, 11, 12, 13 (0-indexed: 4, 7, 9, 10, 11, 12)
    # Since we might delete slides later, we do text edits FIRST before shrinking the presentation.
    
    # --- Slide 5 (Index 4): Make concise ---
    slide_5 = prs.slides[4]
    for shape in slide_5.shapes:
        if not shape.has_text_frame: continue
        text = shape.text
        if "Without East–West Detection" in text:
            concenise_text(shape, "▸ Hidden Hazard  Horizontal sliding precedes vertical collapse.\n▸ Late Warning  Depth-only sensors miss early failure signs.", 12)
        elif "MineGuard — Horizontal First" in text:
            concenise_text(shape, "▸ True E–W Measurement  Isolates pure lateral drift.\n▸ Early Intervention  Provides months of advance warning.", 12)

    # --- Slide 8 (Index 7): Methodology ---
    slide_8 = prs.slides[7]
    for shape in slide_8.shapes:
        if not shape.has_text_frame: continue
        text = shape.text
        if "Free Sentinel-1 C-band SAR images" in text:
            concenise_text(shape, "▸ ESA CDSE OData API retrieval.\n▸ Auto-matching ASC/DSC pairs.", 10)
        elif "Master and slave SLC images" in text:
            concenise_text(shape, "▸ Cross-correlation co-registration.\n▸ Snaphu phase unwrapping.", 10)
        elif "One ASC LOS and one matched DSC" in text:
            concenise_text(shape, "▸ 2x2 linear system solved via Cramer's Rule.\n▸ Isolates true East-West drift.", 10)
        elif "Cumulative displacement is compared" in text:
            concenise_text(shape, "▸ LOW / MED / HIGH / CRITICAL bands.\n▸ Auto-generated action reports.", 10)

    # --- Slide 10 (Index 9): Pages 1 & 2 ---
    slide_10 = prs.slides[9]
    for shape in slide_10.shapes:
        if not shape.has_text_frame: continue
        if "Types the name of the mine" in shape.text:
            concenise_text(shape, "▸ Real-time OpenStreetMap geocoding converts mine name to GPS coordinates in seconds.", 11)
        elif "Clicks four corner points" in shape.text:
            concenise_text(shape, "▸ Interactive map bound-box defines precise ROI, drastically reducing unnecessary data downloads.", 11)

    # --- Slide 11 (Index 10): Pages 3 & 4 ---
    slide_11 = prs.slides[10]
    for shape in slide_11.shapes:
        if not shape.has_text_frame: continue
        if "Selects a Start Year" in shape.text:
            concenise_text(shape, "▸ Users define start/end years and frequency. App provides instant data size estimator.", 11)
        elif "Clicks 'Start Download'" in shape.text:
            concenise_text(shape, "▸ Automated background downloading from ESA Copernicus API with full file integrity checks.", 11)

    # --- Slide 12 (Index 11): Pages 5 & 6 ---
    slide_12 = prs.slides[11]
    for shape in slide_12.shapes:
        if not shape.has_text_frame: continue
        if "The app launches Python workers" in shape.text:
            concenise_text(shape, "▸ Live dashboard orchestrates SNAP InSAR processing with smart queueing and skip-logic.", 11)
        elif "Clicks directly on the satellite map" in shape.text:
            concenise_text(shape, "▸ Boundary-safe interactive markers target exact danger zones without leaving the browser.", 11)

    # --- Slide 13 (Index 12): Pages 7 & 8 ---
    slide_13 = prs.slides[12]
    for shape in slide_13.shapes:
        if not shape.has_text_frame: continue
        if "Cumulative East–West displacement" in shape.text:
            concenise_text(shape, "▸ Automatically plots E/W drift, LOS, and 3D trajectories with auto-generated plain-English findings.", 11)
        elif "Calculates cumulative displacement per hotspot" in shape.text:
            concenise_text(shape, "▸ Calculates threshold-based geotechnical severity and recommends concrete engineering actions.", 11)

    # Shrink presentation to 11 pages by removing 3 less critical slides:
    # We delete indices in reverse order to not shift the earlier ones while deleting
    # 8 = Workflow (redundant with the Page by Page slides)
    # 3 = Use Cases
    # 2 = Design Philosophy
    
    remove_slide(prs, 8)
    remove_slide(prs, 3) 
    remove_slide(prs, 2)

    # Update the slide numbers at the bottom right
    current_count = len(prs.slides)
    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame: continue
            if "/" in shape.text and any(c.isdigit() for c in shape.text) and len(shape.text) < 10:
                # Looks like a page number text box
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                from pptx.enum.text import PP_ALIGN
                p.alignment = PP_ALIGN.RIGHT
                r = p.add_run()
                r.text = f"{i+1} / {current_count}"
                from pptx.util import Pt
                r.font.size = Pt(8)

    prs.save('PD_ppt_edited.pptx')
    print("Done! Edited PD_ppt_edited.pptx")

if __name__ == "__main__":
    main()
